"""Contract tests for extraction, batch isolation, and report integrity."""

import csv
import json
import tempfile
import unittest
from datetime import date
from decimal import Decimal
from pathlib import Path

from openpyxl import load_workbook
from pypdf import PdfWriter

from invoice_analytics.cli import main
from invoice_analytics.extraction import parse_invoice, read_pdf_text
from invoice_analytics.models import BatchResult, Invoice, InvoiceError
from invoice_analytics.pipeline import process_directory
from invoice_analytics.reporting import build_summary, write_reports
from tools.generate_samples import render_pdf

ROOT = Path(__file__).resolve().parents[1]
STANDARD = "Invoice Number: INV-001\nSupplier: Example Supplier\nInvoice Date: 2026-09-01\nCurrency: SGD\nTotal: 125.50"


class ExtractionTests(unittest.TestCase):
    def assert_error(self, text, code):
        with self.assertRaises(InvoiceError) as error:
            parse_invoice(text, "example.pdf")
        self.assertEqual(error.exception.code, code)

    def test_both_layouts_and_whitespace(self):
        standard = parse_invoice("  " + STANDARD.replace("\n", "\r\n  "), "a.pdf")
        vendor_text = STANDARD.replace("Invoice Number:", "Invoice ID:").replace("Supplier:", "Vendor:").replace("Invoice Date:", "Issued:").replace("Total:", "Amount Due:")
        vendor = parse_invoice(vendor_text, "b.pdf")
        self.assertEqual((standard.layout, vendor.layout), ("standard", "vendor"))
        self.assertEqual(standard.total, Decimal("125.50"))
        self.assertEqual(vendor.invoice_date, date(2026, 9, 1))

    def test_missing_repeated_and_unknown_fields(self):
        self.assert_error(STANDARD.replace("Total: 125.50", "Total:"), "missing_field")
        self.assert_error(STANDARD + "\nTotal: 10", "ambiguous_field")
        self.assert_error("An unrelated PDF", "unsupported_layout")
        self.assert_error(STANDARD + "\nInvoice ID: OTHER", "unsupported_layout")

    def test_strict_dates_currency_and_totals(self):
        for value in ["2026-02-30", "01/09/2026", "20260901"]:
            with self.subTest(value=value):
                self.assert_error(STANDARD.replace("2026-09-01", value), "invalid_date")
        self.assert_error(STANDARD.replace("SGD", "JPY"), "unsupported_currency")
        for value in ["-1", "0", "NaN", "Infinity", "1.234", "12,34.50", "$10", "1e3", "1000000000000"]:
            with self.subTest(value=value):
                self.assert_error(STANDARD.replace("125.50", value), "invalid_total")
        self.assertEqual(parse_invoice(STANDARD.replace("125.50", "1,234.50"), "a.pdf").total, Decimal("1234.50"))

    def test_encrypted_pdf_requires_review(self):
        with tempfile.TemporaryDirectory() as temporary:
            path = Path(temporary) / "encrypted.pdf"
            writer = PdfWriter()
            writer.add_blank_page(width=200, height=200)
            writer.encrypt("secret")
            writer.write(path)
            with self.assertRaises(InvoiceError) as error:
                read_pdf_text(path)
            self.assertEqual(error.exception.code, "encrypted_pdf")


class BatchTests(unittest.TestCase):
    def test_demo_matches_reference_records_and_review_reasons(self):
        result = process_directory(ROOT / "data/samples")
        with (ROOT / "data/expected_invoices.csv").open(newline="", encoding="utf-8") as source:
            expected = list(csv.DictReader(source))
        self.assertEqual([invoice.as_row() for invoice in result.invoices], expected)
        with (ROOT / "data/expected_review.csv").open(newline="", encoding="utf-8") as source:
            expected_review = list(csv.DictReader(source))
        self.assertEqual([{"source_file": item.source_file, "code": item.code} for item in result.issues], expected_review)
        self.assertEqual(result.processed, 18)

    def test_conflicting_duplicates_exclude_every_group_member(self):
        for old, new in [("125.50", "999.00"), ("2026-09-01", "2026-09-02"), ("SGD", "USD")]:
            with self.subTest(conflicting_field=old), tempfile.TemporaryDirectory() as temporary:
                folder = Path(temporary)
                render_pdf(folder / "a.pdf", STANDARD.splitlines())
                conflicting = STANDARD.replace(old, new).replace("Example Supplier", "example   supplier")
                render_pdf(folder / "b.PDF", conflicting.splitlines())
                result = process_directory(folder)
                self.assertEqual(result.invoices, [])
                self.assertEqual(
                    [(issue.source_file, issue.code) for issue in result.issues],
                    [("a.pdf", "conflicting_duplicate"), ("b.PDF", "conflicting_duplicate")],
                )
                for issue in result.issues:
                    self.assertIn("a.pdf", issue.message)
                    self.assertIn("b.PDF", issue.message)

    def test_later_conflict_also_excludes_earlier_exact_copies(self):
        with tempfile.TemporaryDirectory() as temporary:
            folder = Path(temporary)
            render_pdf(folder / "a.pdf", STANDARD.splitlines())
            render_pdf(folder / "b.pdf", STANDARD.splitlines())
            render_pdf(folder / "c.pdf", STANDARD.replace("125.50", "999.00").splitlines())
            render_pdf(folder / "d.pdf", STANDARD.replace("INV-001", "INV-002").splitlines())
            result = process_directory(folder)
            self.assertEqual([invoice.source_file for invoice in result.invoices], ["d.pdf"])
            self.assertEqual(
                [(issue.source_file, issue.code) for issue in result.issues],
                [(name, "conflicting_duplicate") for name in ["a.pdf", "b.pdf", "c.pdf"]],
            )
            self.assertEqual(result.processed, 4)

    def test_conflict_outcome_does_not_depend_on_filename_order(self):
        for first_total, last_total in [("125.50", "999.00"), ("999.00", "125.50")]:
            with self.subTest(first_total=first_total), tempfile.TemporaryDirectory() as temporary:
                folder = Path(temporary)
                render_pdf(folder / "a.pdf", STANDARD.replace("125.50", first_total).splitlines())
                render_pdf(folder / "z.pdf", STANDARD.replace("125.50", last_total).splitlines())
                render_pdf(folder / "m.pdf", STANDARD.replace("INV-001", "INV-002").splitlines())
                render_pdf(folder / "n.pdf", STANDARD.replace("Total: 125.50", "Total:").splitlines())
                result = process_directory(folder)
                self.assertEqual([invoice.source_file for invoice in result.invoices], ["m.pdf"])
                self.assertEqual(
                    [(issue.source_file, issue.code) for issue in result.issues],
                    [("a.pdf", "conflicting_duplicate"), ("n.pdf", "missing_field"), ("z.pdf", "conflicting_duplicate")],
                )
                self.assertEqual(build_summary(result)["currencies"]["SGD"]["total"], "125.50")

    def test_exact_duplicates_still_keep_first_file(self):
        with tempfile.TemporaryDirectory() as temporary:
            folder = Path(temporary)
            for name in ["c.pdf", "a.pdf", "b.PDF"]:
                render_pdf(folder / name, STANDARD.splitlines())
            result = process_directory(folder)
            self.assertEqual([invoice.source_file for invoice in result.invoices], ["a.pdf"])
            self.assertEqual(
                [(issue.source_file, issue.code) for issue in result.issues],
                [("b.PDF", "duplicate_invoice"), ("c.pdf", "duplicate_invoice")],
            )

    def test_empty_and_missing_inputs_fail(self):
        with tempfile.TemporaryDirectory() as temporary:
            with self.assertRaises(ValueError):
                process_directory(Path(temporary))
            with self.assertRaises(ValueError):
                process_directory(Path(temporary) / "missing")


class ReportingTests(unittest.TestCase):
    def test_fractional_amounts_sum_without_binary_rounding(self):
        invoices = [
            Invoice(f"{number}.pdf", "standard", str(number), "Supplier", date(2026, 9, 1), "SGD", Decimal(amount))
            for number, amount in enumerate(["0.10", "0.20"])
        ]
        summary = build_summary(BatchResult(invoices, [], 2))
        self.assertEqual(summary["currencies"]["SGD"]["total"], "0.30")

    def test_totals_remain_separate_and_exact(self):
        summary = build_summary(process_directory(ROOT / "data/samples"))
        self.assertEqual(summary["currencies"]["SGD"]["total"], "4196.25")
        self.assertEqual(summary["currencies"]["USD"]["total"], "1175.25")
        self.assertEqual(summary["currencies"]["SGD"]["monthly"], [
            {"month": "2026-07", "count": 3, "total": "1880.00"},
            {"month": "2026-08", "count": 3, "total": "1780.00"},
            {"month": "2026-09", "count": 4, "total": "536.25"},
        ])

    def test_exports_and_reruns(self):
        result = process_directory(ROOT / "data/samples")
        with tempfile.TemporaryDirectory() as temporary:
            folder = Path(temporary) / "reports"
            write_reports(result, folder)
            write_reports(result, folder)
            with (folder / "invoices.csv").open(newline="", encoding="utf-8-sig") as handle:
                self.assertEqual(len(list(csv.DictReader(handle))), 12)
            workbook = load_workbook(folder / "spend_report.xlsx", data_only=False)
            self.assertEqual(workbook.sheetnames, ["Invoices", "Review", "Currency totals", "Monthly spend", "Supplier spend"])
            self.assertEqual(workbook["Invoices"].max_row, 13)
            self.assertEqual(workbook["Review"].max_row, 7)
            self.assertEqual(workbook["Invoices"]["G2"].value, 125.5)
            workbook.close()
            summary = json.loads((folder / "summary.json").read_text())
            self.assertEqual(summary["review_files"], 6)
            html = (folder / "dashboard.html").read_text(encoding="utf-8")
            self.assertIn('data-currency="SGD"', html)
            self.assertIn('data-currency="USD"', html)

    def test_untrusted_text_is_not_a_formula_or_html(self):
        supplier = '=HYPERLINK("https://example.invalid")<script>alert(1)</script>'
        invoice = Invoice("@source.pdf", "standard", "+INV", supplier, date(2026, 9, 1), "SGD", Decimal("10.00"))
        with tempfile.TemporaryDirectory() as temporary:
            folder = Path(temporary)
            write_reports(BatchResult([invoice], [], 1), folder)
            with (folder / "invoices.csv").open(newline="", encoding="utf-8-sig") as handle:
                row = next(csv.DictReader(handle))
                self.assertTrue(row["supplier"].startswith("'="))
                self.assertEqual(row["source_file"], "'@source.pdf")
            workbook = load_workbook(folder / "spend_report.xlsx")
            self.assertEqual(workbook["Invoices"]["D2"].data_type, "s")
            workbook.close()
            html = (folder / "dashboard.html").read_text(encoding="utf-8")
            self.assertNotIn("<script>alert(1)</script>", html)
            self.assertIn("&lt;script&gt;", html)

    def test_all_rejected_batch_has_usable_reports(self):
        with tempfile.TemporaryDirectory() as temporary:
            folder = Path(temporary)
            (folder / "bad.pdf").write_bytes(b"not a PDF")
            result = process_directory(folder)
            write_reports(result, folder / "reports")
            self.assertEqual(len(result.invoices), 0)
            self.assertIn("No invoices passed validation", (folder / "reports/dashboard.html").read_text(encoding="utf-8"))

    def test_cli_exit_statuses(self):
        with tempfile.TemporaryDirectory() as temporary:
            folder = Path(temporary)
            args = ["--input", str(ROOT / "data/samples"), "--output", str(folder / "reports")]
            self.assertEqual(main(args), 0)
            self.assertEqual(main(args + ["--fail-on-review"]), 2)
            self.assertEqual(main(["--input", str(folder / "missing")]), 1)


if __name__ == "__main__":
    unittest.main()
