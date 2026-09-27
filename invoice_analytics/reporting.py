"""Create spreadsheet exports and a self-contained HTML spending dashboard."""

import csv
import json
from collections import defaultdict
from dataclasses import asdict
from decimal import Decimal
from html import escape
from pathlib import Path
from tempfile import TemporaryDirectory

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

from .models import BatchResult


INVOICE_COLUMNS = ["source_file", "layout", "invoice_number", "supplier", "invoice_date", "currency", "total"]
ISSUE_COLUMNS = ["source_file", "code", "message"]
OUTPUT_NAMES = ("invoices.csv", "review.csv", "summary.json", "spend_report.xlsx", "dashboard.html")


def spreadsheet_text(value: str) -> str:
    """Prevent extracted text from being interpreted as a spreadsheet formula."""
    return "'" + value if value.lstrip().startswith(("=", "+", "-", "@")) else value


def write_csv(path: Path, columns: list[str], rows: list[dict]) -> None:
    with path.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=columns)
        writer.writeheader()
        for row in rows:
            writer.writerow({key: spreadsheet_text(str(value)) for key, value in row.items()})


def build_summary(result: BatchResult) -> dict:
    currencies = {}
    for currency in sorted({invoice.currency for invoice in result.invoices}):
        invoices = [invoice for invoice in result.invoices if invoice.currency == currency]
        monthly = defaultdict(lambda: {"count": 0, "total": Decimal("0.00")})
        suppliers = defaultdict(lambda: {"count": 0, "total": Decimal("0.00")})
        names = {}
        for invoice in invoices:
            month = invoice.invoice_date.strftime("%Y-%m")
            supplier_key = invoice.supplier.casefold()
            names.setdefault(supplier_key, invoice.supplier)
            for bucket in (monthly[month], suppliers[supplier_key]):
                bucket["count"] += 1
                bucket["total"] += invoice.total

        currencies[currency] = {
            "invoice_count": len(invoices),
            "supplier_count": len(suppliers),
            "total": format(sum((invoice.total for invoice in invoices), Decimal("0.00")), ".2f"),
            "monthly": [{"month": month, "count": item["count"], "total": format(item["total"], ".2f")} for month, item in sorted(monthly.items())],
            "suppliers": [{"supplier": names[key], "count": item["count"], "total": format(item["total"], ".2f")} for key, item in sorted(suppliers.items(), key=lambda pair: (-pair[1]["total"], pair[0]))],
        }

    return {
        "processed_files": result.processed,
        "accepted_invoices": len(result.invoices),
        "review_files": len(result.issues),
        "currencies": currencies,
    }


def write_workbook(path: Path, result: BatchResult, summary: dict) -> None:
    workbook = Workbook()
    workbook.remove(workbook.active)
    rows = [[invoice.source_file, invoice.layout, invoice.invoice_number, invoice.supplier,
             invoice.invoice_date, invoice.currency, invoice.total] for invoice in result.invoices]
    sheets = {
        "Invoices": (INVOICE_COLUMNS, rows),
        "Review": (ISSUE_COLUMNS, [[item.source_file, item.code, item.message] for item in result.issues]),
        "Currency totals": (["currency", "invoice_count", "supplier_count", "total"], [
            [currency, item["invoice_count"], item["supplier_count"], Decimal(item["total"])]
            for currency, item in summary["currencies"].items()
        ]),
        "Monthly spend": (["currency", "month", "invoice_count", "total"], [
            [currency, item["month"], item["count"], Decimal(item["total"])]
            for currency, data in summary["currencies"].items() for item in data["monthly"]
        ]),
        "Supplier spend": (["currency", "supplier", "invoice_count", "total"], [
            [currency, item["supplier"], item["count"], Decimal(item["total"])]
            for currency, data in summary["currencies"].items() for item in data["suppliers"]
        ]),
    }
    for name, (headers, data) in sheets.items():
        sheet = workbook.create_sheet(name)
        sheet.append(headers)
        for row in data:
            sheet.append(row)
            for cell in sheet[sheet.max_row]:
                if isinstance(cell.value, str):
                    cell.data_type = "s"
                if cell.column == len(headers) and name != "Review":
                    cell.number_format = '#,##0.00'
                if name == "Invoices" and cell.column == 5:
                    cell.number_format = "yyyy-mm-dd"
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        for cell in sheet[1]:
            cell.font = Font(color="FFFFFF", bold=True)
            cell.fill = PatternFill("solid", fgColor="153B4B")
        for column in sheet.columns:
            width = min(75, max(16, max(len(str(cell.value or "")) for cell in column) + 2))
            sheet.column_dimensions[column[0].column_letter].width = width
    workbook.save(path)


def money(value: str) -> str:
    return format(Decimal(value), ",.2f")


def chart_rows(rows: list[dict], label_key: str) -> str:
    maximum = max((Decimal(row["total"]) for row in rows), default=Decimal("1"))
    return "".join(
        f'<div class="bar-row"><div class="bar-label"><span>{escape(row[label_key])}</span>'
        f'<strong>{money(row["total"])}</strong></div><div class="track">'
        f'<div class="bar" style="width:{Decimal(row["total"]) / maximum * 100:.2f}%"></div>'
        '</div></div>' for row in rows
    )


def write_dashboard(path: Path, result: BatchResult, summary: dict, *, reviewed: bool = False) -> None:
    count_label = "Rejected invoices" if reviewed else "Files requiring review"
    queue_title = "Rejected invoices" if reviewed else "Review queue"
    queue_description = (
        "These invoices were rejected during review and are excluded from spending. Decision notes are listed below."
        if reviewed else "Each excluded file has a reason. Correct the source document or verify it manually, then rerun the batch."
    )
    sections = []
    for currency, data in summary["currencies"].items():
        supplier_rows = "".join(
            f'<tr><td>{escape(item["supplier"])}</td><td>{item["count"]}</td><td>{money(item["total"])}</td></tr>'
            for item in data["suppliers"]
        )
        top = data["suppliers"][0]
        share = Decimal(top["total"]) / Decimal(data["total"]) * 100
        sections.append(f'''
        <section class="currency-panel" data-currency="{currency}" aria-label="{currency} spending">
          <div class="section-heading"><h2>{currency} spending</h2><span>Accepted invoices only</span></div>
          <div class="metrics">
            <article><span>Total invoiced</span><strong>{currency} {money(data['total'])}</strong></article>
            <article><span>Invoices</span><strong>{data['invoice_count']}</strong></article>
            <article><span>Suppliers</span><strong>{data['supplier_count']}</strong></article>
          </div>
          <p class="insight">{escape(top['supplier'])} accounts for <strong>{share:.1f}%</strong> of accepted {currency} spending.</p>
          <div class="charts"><article class="card"><h3>Monthly spending</h3><p>Grouped by invoice date · {currency}</p>{chart_rows(data['monthly'], 'month')}</article>
          <article class="card"><h3>Top suppliers</h3><p>Up to 10 suppliers · {currency}</p>{chart_rows(data['suppliers'][:10], 'supplier')}</article></div>
          <article class="card"><h3>Supplier breakdown</h3><div class="table-wrap"><table>
          <thead><tr><th scope="col">Supplier</th><th scope="col">Invoices</th><th scope="col">Total ({currency})</th></tr></thead><tbody>{supplier_rows}</tbody></table></div></article>
        </section>''')

    review_rows = "".join(f'<tr><td>{escape(issue.source_file)}</td><td><code>{escape(issue.code)}</code></td><td>{escape(issue.message)}</td></tr>' for issue in result.issues)
    empty_queue = "No invoices were rejected." if reviewed else "No files require review."
    review_table = f'<div class="table-wrap"><table><thead><tr><th scope="col">File</th><th scope="col">Reason</th><th scope="col">Action / detail</th></tr></thead><tbody>{review_rows}</tbody></table></div>' if review_rows else f'<p>{empty_queue}</p>'
    options = ''.join(f'<option value="{currency}">{currency}</option>' for currency in summary['currencies'])
    document = '''<!doctype html>
<html lang="en"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width, initial-scale=1">
<title>Invoice Analytics | Spend report</title>
<style>
:root{color-scheme:light;--ink:#153b4b;--muted:#506774;--accent:#087f80;--line:#dce6e7}
*{box-sizing:border-box}body{margin:0;background:#f3f6f6;color:var(--ink);font:16px/1.6 system-ui,sans-serif}
header{background:#153b4b;color:#fff;padding:44px max(24px,calc((100% - 1120px)/2))}header p{color:#cce1e4;max-width:680px}
.eyebrow{text-transform:uppercase;letter-spacing:.16em;font-size:12px;font-weight:700}h1{font-size:clamp(30px,5vw,44px);line-height:1.15;margin:12px 0}h2{font-size:24px}h3{margin:0 0 6px}main{max-width:1168px;margin:auto;padding:28px 24px 56px}
.status,.metrics,.charts{display:grid;gap:18px;grid-template-columns:repeat(3,minmax(0,1fr))}.status article,.metrics article,.card{padding:24px;background:#fff;border:1px solid var(--line);border-radius:12px;min-width:0}.status strong,.metrics strong{display:block;font-size:clamp(24px,3vw,34px);margin-top:8px;overflow-wrap:anywhere}article>span,.card>p,.section-heading>span{color:var(--muted);font-size:14px}.review-count{border-top:4px solid #a75c11!important}.toolbar,.section-heading{display:flex;align-items:center;justify-content:space-between;gap:16px;flex-wrap:wrap}.toolbar{margin:28px 0}select{padding:10px 14px;border:1px solid #91a8ac;border-radius:6px;font:inherit;background:white;color:var(--ink)}a{color:#006f70;text-underline-offset:3px}a:focus-visible,select:focus-visible{outline:3px solid #ae6415;outline-offset:4px}.downloads{display:flex;gap:18px;flex-wrap:wrap}.charts{grid-template-columns:repeat(2,minmax(0,1fr));margin:20px 0}.insight{padding:16px 20px;border-left:4px solid var(--accent);background:#e1f1ee}.bar-row{margin-top:20px}.bar-label{display:flex;justify-content:space-between;gap:16px;font-size:14px}.bar-label span{overflow-wrap:anywhere}.track{height:10px;background:#e7eeee;border-radius:5px;margin-top:8px}.bar{height:10px;background:var(--accent);border-radius:5px}.table-wrap{overflow:auto}table{width:100%;border-collapse:collapse;font-size:14px}th,td{text-align:left;padding:14px 12px;border-bottom:1px solid var(--line);vertical-align:top}th{color:var(--muted);font-size:12px;text-transform:uppercase;letter-spacing:.04em}code{font-size:12px}.currency-panel{margin-bottom:32px}.note{color:var(--muted);font-size:14px}.empty{padding:24px;background:white;border:1px solid var(--line)}footer{margin-top:24px}[hidden]{display:none!important}
@media(max-width:700px){.status,.metrics,.charts{grid-template-columns:1fr}header{padding:30px 24px}main{padding:20px 16px}.bar-label{flex-wrap:wrap}}
@media print{header{background:white;color:#153b4b}header p{color:#506774}.toolbar{display:none}.currency-panel[hidden]{display:block!important}.card{break-inside:avoid}}
</style></head><body><header><div class="eyebrow">Invoice Analytics / Operations report</div><h1>From invoices to spending insights.</h1><p>Validated invoice records, supplier spending, and a review queue for documents that need attention.</p></header><main>
'''
    document += f'''<div class="status"><article><span>PDFs processed</span><strong>{result.processed}</strong></article><article><span>Accepted invoices</span><strong>{len(result.invoices)}</strong></article><article class="review-count"><span>{count_label}</span><strong>{len(result.issues)}</strong></article></div>
<div class="toolbar"><label for="currency">Currency <select id="currency"><option value="all">All currencies</option>{options}</select></label><nav class="downloads" aria-label="Report downloads"><a href="invoices.csv">Invoice CSV</a><a href="spend_report.xlsx">Excel workbook</a><a href="review.csv">Review CSV</a></nav></div>
<p class="note">Currencies are reported separately. Spending reflects invoiced amounts, not verified payments. Excluded files do not contribute to totals. Months without accepted invoices are omitted.</p>'''
    empty_message = "No invoices were approved. All records in this report were rejected." if reviewed else "No invoices passed validation. Review the files below before using this batch for analysis."
    document += ''.join(sections) or f'<p class="empty">{empty_message}</p>'
    document += f'<article class="card"><h2>{queue_title}</h2><p>{queue_description}</p>{review_table}</article>'
    document += '''<footer class="note">Generated locally by Invoice Analytics Automation. No external services or exchange-rate conversion are used.</footer></main>
<script>document.getElementById('currency').addEventListener('change',function(){for(const panel of document.querySelectorAll('.currency-panel')){panel.hidden=this.value!=='all'&&panel.dataset.currency!==this.value;}});</script>
</body></html>'''
    path.write_text(document, encoding="utf-8")


def write_reports(result: BatchResult, output_dir: Path, *, reviewed: bool = False) -> dict:
    summary = build_summary(result)
    output_dir.parent.mkdir(parents=True, exist_ok=True)
    # Finish generating all files before replacing reports from a previous run.
    with TemporaryDirectory(prefix="invoice-report-", dir=output_dir.parent) as temporary:
        stage = Path(temporary)
        write_csv(stage / "invoices.csv", INVOICE_COLUMNS, [invoice.as_row() for invoice in result.invoices])
        write_csv(stage / "review.csv", ISSUE_COLUMNS, [asdict(issue) for issue in result.issues])
        (stage / "summary.json").write_text(json.dumps(summary, indent=2) + "\n", encoding="utf-8")
        write_workbook(stage / "spend_report.xlsx", result, summary)
        write_dashboard(stage / "dashboard.html", result, summary, reviewed=reviewed)
        output_dir.mkdir(parents=True, exist_ok=True)
        for name in OUTPUT_NAMES:
            (stage / name).replace(output_dir / name)
    return summary
